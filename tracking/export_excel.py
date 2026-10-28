"""Regenerate both team Excel trackers from the repo tracker (the master copy).

    python tracking/export_excel.py

Sources (never edited): tracking/exports/source/*.xlsm, the original Drive files.
Outputs: tracking/exports/QuantEdge_1_0_Project_Management_Tracker.xlsm
         tracking/exports/QuantEdge_FigJam_Pipeline_Block_Reference.xlsm
Original subtask and task wording, owners and weeks are kept for the 98 original subtasks; status, priority,
outputs, dependencies and new subtasks come from SUBTASKS.md. Formulas and dropdown ranges are widened to row LAST.
Upload each file to Drive with "Manage versions → Upload new version" to keep the same link.
"""
from __future__ import annotations

import copy
import re
from pathlib import Path

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
SRC, OUT = HERE / "exports" / "source", HERE / "exports"
PM, BR = "QuantEdge_1_0_Project_Management_Tracker.xlsm", "QuantEdge_FigJam_Pipeline_Block_Reference.xlsm"
LAST = 400
STATUS = {"Not started": "Not Started", "In progress": "In Progress", "Done": "Done"}
OWNERS = '"Malindu,Raveena,Dewmini,Sujan,Raveena + Dewmini,All members"'
WEEKS = '"Week 1,Week 2,Week 3,Weeks 1–2,Weeks 2–3,Ongoing,If finalist"'
STATES = '"Not Started,In Progress,Blocked,In Review,Done"'


def parse_tracker() -> tuple[dict, list[list[str]], dict]:
    """Tasks (from ## headers), subtask rows and definition-of-done per task from the markdown tracker."""
    text = (HERE / "SUBTASKS.md").read_text(encoding="utf-8")
    tasks, rows = {}, []
    for m in re.finditer(r"^## (T\d+) — (.+)\nStage: (.*?) · Owner: (.*?) · Priority: (.*?) · Target: (.*)$", text, re.M):
        tasks[m.group(1)] = {"name": m.group(2), "stage": m.group(3), "owner": m.group(4), "priority": m.group(5), "week": m.group(6)}
    for line in text.splitlines():
        if re.match(r"^\| T\d+\.\d+ \|", line):
            rows.append([c.strip() for c in line.strip().strip("|").split("|")])
    dod = {}
    section = (HERE / "TASKS.md").read_text(encoding="utf-8").split("## Verification criteria", 1)[1].split("\n## ", 1)[0]
    for m in re.finditer(r"^\| T(\d+)(?:–T(\d+))? \| (.+) \|$", section, re.M):
        for i in range(int(m.group(1)), int(m.group(2) or m.group(1)) + 1):
            dod[f"T{i}"] = m.group(3)
    return tasks, rows, dod


def wb_status(s: str) -> str:
    return "Blocked" if s.startswith("Blocked") else STATUS[s]


def task_status(statuses: list[str]) -> str:
    if all(s == "Done" for s in statuses):
        return "Done"
    if any(s == "Blocked" for s in statuses):
        return "Blocked"
    return "In Progress" if any(s in ("Done", "In Progress") for s in statuses) else "Not Started"


def styles(ws, row: int, ncol: int) -> list:
    return [copy.copy(ws.cell(row, c)._style) for c in range(1, ncol + 1)]


def reset_validation(ws, specs: list[tuple[str, str]]) -> None:
    ws.data_validations.dataValidation = []
    for rng, formula in specs:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.add(rng)
        ws.add_data_validation(dv)


def reset_cf(ws, rng: str) -> None:
    old = list(ws.conditional_formatting)
    ws.conditional_formatting = type(ws.conditional_formatting)()
    for cf in old:
        for rule in cf.rules:
            ws.conditional_formatting.add(rng, rule)


def export_pm(tasks: dict, rows: list[list[str]], dod: dict) -> Path:
    wb = openpyxl.load_workbook(SRC / PM, keep_vba=True)
    mt, st, db = wb["Main Tasks"], wb["Subtasks"], wb["Dashboard"]
    orig_sub = {st.cell(r, 1).value: [st.cell(r, c).value for c in range(1, 11)] for r in range(5, st.max_row + 1) if st.cell(r, 1).value}
    orig_task = {mt.cell(r, 1).value: {c: mt.cell(r, c).value for c in range(1, 12)} for r in range(5, mt.max_row + 1) if mt.cell(r, 1).value}

    # Subtasks sheet
    sub_style, hdr = styles(st, 5, 10), copy.copy(st.cell(4, 10)._style)
    st.delete_rows(5, st.max_row)
    for c, name in enumerate(["Expected Output", "Priority", "Depends On", "Source"], 11):
        st.cell(4, c, name)._style = hdr
    for col, w in zip("KLMN", [44, 12, 22, 30]):
        st.column_dimensions[col].width = w
    by_task: dict[str, list[str]] = {}
    for i, (sid, text, output, pri, owner, dep, status, source) in enumerate(rows):
        task, rr = sid.split(".")[0], 5 + i
        o = orig_sub.get(sid)
        s = wb_status(status)
        by_task.setdefault(task, []).append(s)
        vals = [sid, task, orig_task[task][2] if task in orig_task else tasks[task]["name"],
                o[3] if o else text, o[4] if o else owner.replace(" (coordinates)", ""), s,
                o[6] if o else tasks[task]["week"], o[7] if o else None,
                output if s == "Done" else None, o[9] if o else None, output, pri, dep, source]
        for c, v in enumerate(vals, 1):
            st.cell(rr, c, v)._style = copy.copy(sub_style[min(c, 10) - 1])
        st.row_dimensions[rr].height = 36.75
    st.cell(2, 1).value = (f"Mark individual subtasks Done; main-task progress updates automatically. {len(rows)} subtasks "
                           "(98 original + added). Generated from tracking/SUBTASKS.md by tracking/export_excel.py.")
    reset_validation(st, [(f"E5:E{LAST}", OWNERS), (f"F5:F{LAST}", STATES), (f"G5:G{LAST}", WEEKS),
                          (f"L5:L{LAST}", '"Critical,High,Medium,Low"')])
    reset_cf(st, f"F5:F{LAST}")

    # Main Tasks sheet
    top, rest = styles(mt, 5, 11), styles(mt, 6, 11)
    for m in list(mt.merged_cells.ranges):
        if m.min_row >= 5:
            mt.unmerge_cells(str(m))
    mt.delete_rows(5, mt.max_row)
    mt.cell(4, 12, "Definition of Done")._style = copy.copy(mt.cell(4, 11)._style)
    mt.column_dimensions["L"].width = 60
    r = 5
    for task, meta in tasks.items():
        subs = [x for x in rows if x[0].split(".")[0] == task]
        o = orig_task.get(task)
        values = {1: task, 2: o[2] if o else meta["name"], 3: o[3] if o else meta["name"],
                  5: o[5] if o else meta["owner"].replace(" (coordinates)", ""),
                  6: o[6] if o else ("Core ★" if meta["priority"] == "Core" else meta["priority"]),
                  7: task_status(by_task[task]),
                  8: f'=IFERROR(COUNTIFS(Subtasks!$B$5:$B${LAST},A{r},Subtasks!$F$5:$F${LAST},"Done")/COUNTIF(Subtasks!$B$5:$B${LAST},A{r}),0)',
                  9: o[9] if o else meta["week"], 10: o[10] if o else None, 11: o[11] if o else None, 12: dod.get(task)}
        for i, x in enumerate(subs):
            for c in range(1, 13):
                mt.cell(r + i, c)._style = copy.copy((top if i == 0 else rest)[min(c, 11) - 1])
            mt.cell(r + i, 4, f"{x[0].split('.')[1]}  {orig_sub[x[0]][3] if x[0] in orig_sub else x[1]}")
            mt.row_dimensions[r + i].height = 34.5
        for c, v in values.items():
            mt.cell(r, c, v)
        if len(subs) > 1:
            for c in [1, 2, 3, 5, 6, 7, 8, 9, 10, 11, 12]:
                mt.merge_cells(start_row=r, start_column=c, end_row=r + len(subs) - 1, end_column=c)
        r += len(subs)
    reset_validation(mt, [(f"F5:F{LAST}", '"Core ★,Supporting,Optional,Conditional"'), (f"E5:E{LAST}", OWNERS),
                          (f"G5:G{LAST}", STATES), (f"I5:I{LAST}", WEEKS)])
    reset_cf(mt, f"G5:G{LAST}")

    # Dashboard: widen ranges, headline, review list
    for row in db.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = c.value.replace("$102", f"${LAST}")
    db["A2"] = f"{len(tasks)} tasks | {len(rows)} subtasks | status and subtask progress"
    rs = styles(db, 28, 4)
    for rr, (tid, owner) in zip([29, 30], [("T20", "All members"), ("T21", "Malindu")]):
        for c, v in enumerate([tid, tasks[tid]["name"], owner, None], 1):
            db.cell(rr, c, v)._style = copy.copy(rs[c - 1])
    for row in wb["How to Use"].iter_rows():
        for c in row:
            if isinstance(c.value, str) and "lists all 19 tasks" in c.value:
                c.value = c.value.replace("lists all 19 tasks", f"lists all {len(tasks)} tasks; column L = definition of done")
    path = OUT / PM
    wb.save(path)
    return path


CHANGES = [  # (pipeline, block, previous, new, why, source, decision)
    ("Master Workflow", "02", "Client rules / policy (cash needs)", "Client rules: spending 4% paid by selling AGG; no cash sleeve", "Brief allows only the 7 ETFs", "pp.4–5", "D10"),
    ("Mandate & Data", "01", "Return objective & real-growth hurdle", "Required return (1+I)(1+S)(1+G)(1+C): target 7.19%, minimum 6.66%; fallback = highest sustainable spending", "Needs-based, risk-constrained objective", "pp.3–4, p.8", "D12"),
    ("Mandate & Data", "02", "Risk & liquidity budget", "DD 25% (P ≤ 5% in 5y), ES99 1-day ≤ 2.5% NAV, stress ≤ 25%, AGG ≥ 8%", "Probabilistic, tiered limits; research doc §2/§4b", "pp.3–4", "D7"),
    ("Mandate & Data", "06", "Train / test governance", "Train 2004–14 / validation 2015–19 / test 2020–26", "More validation power; unseen regime in test", "p.7", "D5"),
    ("Portfolio Models", "NEW", "—", "Minimum-CVaR (target return)", "Optimises the client's downside measure directly", "p.7 (team choice)", "D20"),
    ("Portfolio Models", "06", "Feasibility filters", "D6 v2: equity 20–45%, GLD ≤ 20%, each ≤ 2× benchmark, XLE/XLK ≤ 10% each / 15% together, component-ES cap", "Derived from drawdown vs purchasing-power risk", "p.4", "D6"),
    ("Portfolio Models", "N01", "CAPM and benchmark exposure", "SPY = CAPM proxy; Policy 40/50/10 primary; 60/40 and EW7 secondary; ERP 5%, monthly betas", "60/40 fails 'appropriate' for a conservative mandate", "p.7", "D3, D12"),
    ("Portfolio Models", "08", "P1 handoff", "P1-train (fit 2004–14, for OOS tests) and P1-final (full-sample re-fit, same method)", "Avoids leakage while using all data for the final decision", "p.7", "D5"),
    ("Risk Estimation", "N01", "Confidence / horizon control", "90/95/99%, 1-day primary, uncertainty bands on every estimate", "Brief asks 'how confident'", "p.5", "D17, D18"),
    ("Risk Estimation", "04", "Monte Carlo VaR / ES", "Multivariate Student-t / t-copula, fixed seed", "Normal MC would duplicate parametric VaR", "p.7", "D19"),
    ("Validation & Decision", "08 / 10", "Accept • resize • reject", "Accept the mandate as proposed • adjust the risk limit • reject", "Brief's exact wording", "p.8", "D14"),
    ("Validation & Decision", "09", "Revise & retest (hedge/cash buffer)", "Revise target, constraints or model; no hedge or cash buffer", "No derivatives or cash in mandate", "pp.4–5", "D10"),
]
CHECK_STATUS = {  # Official Checklist category -> (status 2026-10-28, evidence / decision)
    "Client": ("Decided, IPS doc pending", "D4 (proposed), D6, D7, D12 approved"),
    "Universe": ("Met", "validate_weights rejects other ETFs (tested)"),
    "Portfolio methods": ("Partly built", "CAPM betas/expected returns built; Markowitz pending (T4)"),
    "KPIs": ("Not started", "T5; formulas in CALCULATION_DICTIONARY"),
    "VaR trio": ("Not started", "T6; MC distribution D19"), "Expected Shortfall": ("Not started", "T7"),
    "EVT": ("Not started", "T8; threshold method D8 proposed"), "Risk contribution": ("Not started", "T9; D6 layer 4"),
    "Backtesting": ("Not started", "T10; split D5 approved"), "Crisis scenarios": ("Partly built", "Windows used in research; hypothetical D9 open"),
    "Data governance": ("Met (sign-off pending)", "src/quantedge/data, data/README.md, tests"),
    "Streamlit interface": ("Not started", "T14"), "Limitations": ("Doc drafted, app pending", "docs/ASSUMPTIONS_LIMITATIONS.md"),
    "Deliverables": ("Partly", "run_all.py works for data + portfolio; no commit/report yet"),
    "AI governance": ("Draft", "docs/AI_DISCLOSURE.md draft; 9 validation rows unsigned"),
    "Optional extensions": ("CORRECTION", "Brief p.10 names regime-conditioned risk, risk overlays, alternative EVT thresholds; copulas/wavelets are team extensions, not brief items"),
    "Judging weighting": ("Info", "Decision 30%, Technical 25%, Software/Innovation/Communication 15% each"),
    "Cost and real growth": ("Decided", "D12: CPI via 10y breakeven, S 4%, G 0.5%, C 0.2%; ETF fees not double-counted"),
}
NEW_CHECKS = [
    ("Trade-offs & diversification", "Quantify return vs risk-reduction trade-offs; show diversification benefits", "Portfolio Models 07 / N03", "p.4", "REQUIRED", "Partly", "Drawdown↔spending table (research §4); T4.14"),
    ("Confidence in estimates", "State how confident we are in risk estimates (uncertainty bands)", "Risk Estimation 08", "p.5", "REQUIRED", "Decided (D18)", "T6.09, T8.09"),
    ("Models under stress", "Show how risk models perform under stressed conditions", "Risk Estimation 06 / Validation 03", "p.7", "REQUIRED", "Not started", "T8.11"),
    ("Assumption critique", "Evaluate and justify assumptions behind each risk model", "Risk Estimation 02–04", "p.7", "REQUIRED", "Decided (D21)", "T6.08"),
    ("App confidence 90%", "Sidebar offers 90/95/99%; every measure available at 90%", "App 02 / Risk N01", "p.5", "REQUIRED", "Configured", "risk.confidence_levels; T6.07"),
]


def export_br() -> Path:
    wb = openpyxl.load_workbook(SRC / BR, keep_vba=True)
    cr = wb["Change Register"]
    style, start = styles(cr, cr.max_row, 8), cr.max_row + 1
    for i, (pipe, block, prev, new, why, src, dec) in enumerate(CHANGES):
        for c, v in enumerate(["DECIDED 2026-10-28", pipe, block, prev, new, f"{why} ({dec})", src, "Pending FigJam edit (review §4)"], 1):
            cr.cell(start + i, c, v)._style = copy.copy(style[c - 1])
    ck = wb["Official Checklist"]
    hdr = copy.copy(ck.cell(1, 5)._style)
    for c, name in [(6, "Status 2026-10-28"), (7, "Evidence / decision")]:
        ck.cell(1, c, name)._style = hdr
    ck.column_dimensions["F"].width, ck.column_dimensions["G"].width = 24, 70
    body = styles(ck, 2, 5)
    for r in range(2, ck.max_row + 1):
        status = CHECK_STATUS.get(ck.cell(r, 1).value)
        if status:
            for c, v in [(6, status[0]), (7, status[1])]:
                ck.cell(r, c, v)._style = copy.copy(body[4])
    start = ck.max_row + 1
    for i, row in enumerate(NEW_CHECKS):
        for c, v in enumerate(row, 1):
            ck.cell(start + i, c, v)._style = copy.copy(body[min(c, 5) - 1])
    sh = wb["Start Here"]
    note_style = copy.copy(sh.cell(37, 1)._style)
    sh.cell(sh.max_row + 2, 1, "UPDATE 2026-10-28: live board 0BsezYr2IQNrweujKm3yF3 read and verified. Decisions D3–D22 are "
            "recorded in the repo (docs/DECISIONS.md). 'DECIDED 2026-10-28' rows in Change Register list the FigJam edits still "
            "to apply. Official Checklist columns F–G show current status; 5 brief requirements added.")._style = note_style
    path = OUT / BR
    wb.save(path)
    return path


if __name__ == "__main__":
    tasks, rows, dod = parse_tracker()
    print(export_pm(tasks, rows, dod).name, f"({len(tasks)} tasks, {len(rows)} subtasks)")
    print(export_br().name, f"({len(CHANGES)} change rows, {len(NEW_CHECKS)} new checklist rows)")
